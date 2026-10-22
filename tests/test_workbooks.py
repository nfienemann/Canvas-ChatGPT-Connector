import io
import json
import random
import re
import string
import subprocess
import sys
import zipfile
from datetime import datetime
from pathlib import Path

import httpx
import pytest
from defusedxml.common import DefusedXmlException
from mcp import ClientSession
from mcp.client.streamable_http import streamable_http_client
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from canvas_bridge.course_files import list_files, read_file, search_files, status
from canvas_bridge.domain import DomainError
from canvas_bridge.parser_sandbox import extract_isolated
from canvas_bridge.parser_worker import extract
from canvas_bridge.workbook_worker import extract_workbook
from canvas_bridge.workbooks import read_workbook
from tests.conftest import seed
from tests.test_course_files import file_record
from tests.test_mcp_contract import running as running

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_standalone_import_does_not_open_host_mime_configuration(tmp_path):
    from canvas_bridge import workbook_worker

    script = tmp_path / "import_probe.py"
    script.write_text(
        """import builtins, mimetypes, os.path, runpy, sys
from pathlib import Path
source = Path(sys.argv[1])
sys.path.insert(0, str(source.parent))
host = "/host-private-mime.types"
real_open, real_isfile = builtins.open, os.path.isfile
def blocked_open(path, *args, **kwargs):
    if str(path) == host:
        raise PermissionError("Synthetic host-config denial")
    return real_open(path, *args, **kwargs)
builtins.open = blocked_open
os.path.isfile = lambda path: True if str(path) == host else real_isfile(path)
mimetypes.knownfiles, mimetypes.inited, mimetypes._db = [host], False, None
try:
    mimetypes.MimeTypes()
except PermissionError:
    pass
else:
    raise AssertionError("Host-config denial control did not run")
mimetypes.inited, mimetypes._db = False, None
runpy.run_path(str(source), run_name="__main__")
assert mimetypes.knownfiles == []
print("Standalone workbook import uses no host MIME file")
""",
        encoding="utf-8",
    )
    result = subprocess.run(  # noqa: S603 - fixed synthetic probe and trusted local Python.
        [sys.executable, "-I", str(script), str(Path(workbook_worker.__file__))],
        capture_output=True,
        text=True,
        timeout=20,
        check=True,
    )
    assert "uses no host MIME file" in result.stdout


def workbook():
    wb = Workbook()
    s = wb.active
    s.title = "Analysis"
    s.append(["Sample", "Ct Target", "Ct Reference", "Delta Ct"])
    s.append(["0007", 24.5, 19.2, "=B2-C2"])
    s.append(["Treatment", 22.1, 19.0, "=2^(-D2)"])
    s["A5"] = "Merged header"
    s.merge_cells("A5:D5")
    s["A6"] = datetime(2026, 10, 21)
    s["B6"] = 0.25
    s["B6"].number_format = "0.0%"
    s["C6"] = 1.2e-8
    s["D6"] = "#DIV/0!"
    s["A7"] = True
    s.add_table(Table(displayName="Measurements", ref="A1:D3"))
    hidden = wb.create_sheet("Reference")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Control"
    s["D7"] = "=Reference!A1"
    s["D8"] = "=AVERAGE(InputValues)"
    wb.defined_names.add(DefinedName("InputValues", attr_text="'Analysis'!B2:B3"))
    wb.create_sheet("Empty")
    buf = io.BytesIO()
    wb.save(buf)
    # Add one known stored cached value without calculating it in the reader.
    return rewrite(
        buf.getvalue(),
        "xl/worksheets/sheet1.xml",
        lambda s: re.sub(r"<f>B2-C2</f><v(?:\s*/>|></v>)", "<f>B2-C2</f><v>5.3</v>", s),
    )


def rewrite(data, name, transform):
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dest,
    ):
        for n in source.namelist():
            content = source.read(n)
            dest.writestr(n, transform(content.decode()).encode() if n == name else content)
    return out.getvalue()


def test_T11_T12_structured_xlsx_values_formulas_and_sources():
    metadata, pages, partial, _ = extract_workbook(workbook())
    assert not partial
    assert metadata["worksheet_count"] == 3 and metadata["date_system"] == "1900"
    assert metadata["worksheets"][1]["visibility"] == "hidden"
    assert metadata["worksheets"][2]["used_range"] is None
    assert metadata["worksheets"][0]["merged_regions"] == ["A5:D5"]
    assert metadata["worksheets"][0]["tables"][0]["headers"][0] == "Sample"
    assert metadata["named_ranges"][0]["name"] == "InputValues"
    cells = {c["address"]: c for p in pages if p["sheet"] == "Analysis" for c in p["cells"]}
    assert cells["A2"]["value"] == "0007"
    assert cells["D2"]["formula"] == "=B2-C2" and cells["D2"]["cached_value"] == 5.3
    assert cells["D3"]["cached_value_status"] == "unavailable"
    assert cells["D7"]["cross_sheet_reference"]
    assert cells["B6"]["value"] == 0.25 and cells["B6"]["number_format"] == "0.0%"
    assert cells["A6"]["date_time_interpretation"].startswith("2026-10-21")
    assert cells["A6"]["raw_value"].isdigit()
    assert cells["C6"]["value"] == 1.2e-8 and cells["D6"]["error"] == "#DIV/0!"
    assert cells["A7"]["value"] is True


def test_T11_xlsx_isolated_worker(tmp_path):
    r = extract_isolated(workbook(), "xlsx", tmp_path / "parser")
    assert r["metadata"]["worksheet_count"] == 3
    assert any(c["formula"] == "=B2-C2" for p in r["pages"] for c in p["cells"])


def test_T11_sparse_dimensions_and_wrong_declared_range_are_partial():
    data = rewrite(
        workbook(),
        "xl/worksheets/sheet1.xml",
        lambda s: s.replace('ref="A1:D8"', 'ref="A1:A1"').replace(
            "</sheetData>", '<row r="1000000"><c r="A1000000" t="n"><v>42</v></c></row></sheetData>'
        ),
    )
    metadata, pages, partial, _ = extract_workbook(data)
    assert partial and metadata["worksheets"][0]["used_range"].endswith("1000000")
    assert any(c["address"] == "D8" for p in pages for c in p["cells"])
    assert all(c["row"] <= 5000 for p in pages for c in p["cells"])


@pytest.mark.parametrize(
    "kind", ["corrupt", "traversal", "macro", "entity", "encrypted", "oversized"]
)
def test_T13_malicious_workbooks_rejected(kind):
    data = workbook()
    if kind == "corrupt":
        data = b"PK\x03\x04broken"
    elif kind == "encrypted":
        data = b"\xd0\xcf\x11\xe0EncryptedPackage"
    elif kind == "entity":
        data = rewrite(
            data,
            "xl/workbook.xml",
            lambda s: '<!DOCTYPE x [<!ENTITY x SYSTEM "file:///secret">]>' + s,
        )
    else:
        buf = io.BytesIO(data)
        with zipfile.ZipFile(buf, "a", zipfile.ZIP_DEFLATED) as z:
            z.writestr(
                {
                    "traversal": "../escape.xml",
                    "macro": "xl/vbaProject.bin",
                    "oversized": "large.xml",
                }[kind],
                b"x" * 21_000_000 if kind == "oversized" else b"unsafe",
            )
        data = buf.getvalue()
    with pytest.raises((ValueError, zipfile.BadZipFile, DefusedXmlException)):
        extract_workbook(data)


@pytest.fixture
async def readable_workbook(bridge, owner, seeded, monkeypatch):
    data = workbook()
    bridge.synthetic.data["files"] = {
        "11": [
            file_record(
                filename="fictional.xlsx",
                display_name="fictional.xlsx",
                size=len(data),
                updated_at="2026-10-21T12:00:00Z",
                url="https://canvas.synthetic.invalid/files/51/download",
                **{"content-type": MIME},
            )
        ]
    }
    bridge.synthetic.data["file_bodies"] = {"51": data}
    bridge.synthetic.data["file_types"] = {"51": MIME}
    bridge.settings.file_reading_enabled = True

    async def parse(data, kind, settings):
        return extract(data, kind)

    monkeypatch.setattr("canvas_bridge.course_files.extract_document", parse)
    _, courses = bridge.selected(owner, seeded)
    listed = await list_files(bridge, owner, seeded, courses[0].id)
    return courses[0], listed.data["files"][0]["file_id"]


async def test_T12_T25_xlsx_tools_continuation_search_and_citations(
    bridge, owner, seeded, readable_workbook
):
    course, ref = readable_workbook
    r = await read_workbook(bridge, owner, seeded, ref)
    assert r.data["workbook"]["worksheet_count"] == 3 and not r.data["cells"]
    first = await read_workbook(bridge, owner, seeded, ref, "Analysis", "A1:D3", max_cells=5)
    assert len(first.data["cells"]) == 5 and first.data["continuation"]
    second = await read_workbook(bridge, owner, seeded, **first.data["continuation"])
    assert not {c["address"] for c in first.data["cells"]} & {
        c["address"] for c in second.data["cells"]
    }
    assert second.data["cells"][0]["source_reference"]["locator"].startswith("'Analysis'!")
    assert (await status(bridge, owner, seeded, ref)).data["indexed"]
    search = await search_files(bridge, owner, seeded, course.id, "Treatment")
    assert search.data["results"]
    assert "'Analysis'!" in search.data["results"][0]["passage"]
    assert (await read_file(bridge, owner, seeded, ref)).data["format"] == "xlsx"


@pytest.mark.parametrize(
    "sheet,cell_range",
    [
        ("Missing", None),
        ("Analysis", "D5:A1"),
        ("Analysis", "A0:F5"),
        (None, "A1:B2"),
        ("Analysis", "XFE1"),
    ],
)
async def test_T11_invalid_sheet_range_is_explicit(
    bridge, owner, seeded, readable_workbook, sheet, cell_range
):
    _, ref = readable_workbook
    with pytest.raises(DomainError):
        await read_workbook(bridge, owner, seeded, ref, sheet, cell_range)


async def test_T13_xlsx_cross_user_and_source_change(bridge, owner, seeded, readable_workbook):
    _, ref = readable_workbook
    r = await read_workbook(bridge, owner, seeded, ref, "Analysis", max_cells=1)
    other = await seed(bridge, "other-owner", "other-connection")
    with pytest.raises(DomainError):
        await read_workbook(bridge, "other-owner", other, ref)
    bridge.synthetic.data["files"]["11"][0]["updated_at"] = "2026-10-22T00:00:00Z"
    # Metadata renewal with identical bytes preserves the content hash; changed bytes do not.
    bridge.synthetic.data["file_bodies"]["51"] = rewrite(
        workbook(), "xl/worksheets/sheet1.xml", lambda s: s.replace("24.5", "25.5")
    )
    bridge.synthetic.data["files"]["11"][0]["size"] = len(
        bridge.synthetic.data["file_bodies"]["51"]
    )
    with pytest.raises(DomainError):
        await read_workbook(bridge, owner, seeded, **r.data["continuation"])


def test_dates_1904_cached_errors_external_and_array_formulas():
    data = rewrite(
        workbook(),
        "xl/workbook.xml",
        lambda s: re.sub(r"<workbookPr(?=[\s/>])", '<workbookPr date1904="1"', s),
    )
    metadata, pages, _, _ = extract_workbook(data)
    assert metadata["date_system"] == "1904"
    dates = [c for p in pages for c in p["cells"] if c["address"] == "A6"]
    assert dates[0]["date_time_interpretation"].startswith("2030")
    data = rewrite(
        workbook(),
        "xl/worksheets/sheet1.xml",
        lambda s: (
            s.replace(
                '<c r="D2"><f>B2-C2</f><v>5.3</v>', '<c r="D2" t="e"><f>B2-C2</f><v>#DIV/0!</v>'
            )
            .replace("<f>2^(-D2)</f>", '<f t="array" ref="D3:D3">2^(-D2)</f>')
            .replace("<f>Reference!A1</f>", "<f>'[1]External'!A1</f>")
        ),
    )
    _, pages, _, _ = extract_workbook(data)
    cells = {c["address"]: c for p in pages if p["sheet"] == "Analysis" for c in p["cells"]}
    assert cells["D2"]["error"] == "#DIV/0!" and cells["D2"]["formula"] == "=B2-C2"
    assert cells["D3"]["formula"]["status"] == "unsupported_formula_type"
    assert cells["D3"]["formula"]["text"] == "2^(-D2)"
    assert cells["D7"]["possible_external_reference"]
    assert cells["D7"]["cached_value_status"] == "unavailable"


def test_workbook_bounds_report_partial_not_empty_and_do_not_overflow():
    wb = Workbook()
    rng = random.Random(7)
    for row in range(1, 101):
        wb.active.cell(row, 1, "".join(rng.choices(string.ascii_letters, k=4000)))
    wb.create_sheet("Later")["A1"] = "Still reported after earlier sheet truncation"
    stream = io.BytesIO()
    wb.save(stream)
    result = extract(stream.getvalue(), "xlsx")
    assert result["coverage"] == "partial"
    sheet = result["metadata"]["worksheets"][0]
    assert 0 < sheet["extracted_cell_count"] < sheet["nonempty_cell_count"] == 100
    assert sheet["coverage"] == "partial"
    assert len(result["metadata"]["worksheets"]) == 2
    assert len(json.dumps(result, ensure_ascii=True)) < 2_000_000
    assert sum(len(p["text"]) for p in result["pages"]) < 250_000
    assert not any("index_text" in c for p in result["pages"] for c in p["cells"])


@pytest.mark.parametrize(
    "case", ["sheet_limit", "metadata_limit", "cell_limit", "external_sheet", "wrong_type"]
)
def test_workbook_explicit_failures(case):
    data = workbook()
    expected = {
        "sheet_limit": "workbook_sheet_limit",
        "metadata_limit": "workbook_metadata_limit",
        "cell_limit": "workbook_cell_limit",
        "external_sheet": "invalid_office_reference",
        "wrong_type": "office_type_mismatch",
    }[case]
    if case == "sheet_limit":
        data = rewrite(
            data,
            "xl/workbook.xml",
            lambda s: s.replace("</sheets>", '<sheet name="overflow"/>' * 51 + "</sheets>"),
        )
    elif case == "metadata_limit":
        data = rewrite(data, "xl/workbook.xml", lambda s: s.replace("InputValues", "I" * 31000))
    elif case == "cell_limit":
        data = rewrite(
            data, "xl/worksheets/sheet1.xml", lambda s: s.replace("Treatment", "T" * 8193)
        )
    elif case == "external_sheet":
        data = rewrite(
            data,
            "xl/_rels/workbook.xml.rels",
            lambda s: s.replace(
                'Target="/xl/worksheets/sheet1.xml"',
                'TargetMode="External" Target="https://private.invalid/sheet.xml"',
            ),
        )
    else:
        from tests.test_academic_intelligence import ooxml

        data = ooxml("docx")
    with pytest.raises(ValueError, match=expected):
        extract_workbook(data)


async def test_permission_denial_clears_cached_cells_and_reads_do_not_write_canvas(
    bridge, owner, seeded, readable_workbook
):
    _, ref = readable_workbook
    await read_workbook(bridge, owner, seeded, ref)
    bridge.synthetic.data["files"]["11"][0]["hidden_for_user"] = True
    with pytest.raises(DomainError):
        await read_workbook(bridge, owner, seeded, ref, "Analysis")
    assert not (await status(bridge, owner, seeded, ref)).data["indexed"]
    assert all(r.method == "GET" for r in bridge.synthetic.requests)


async def test_real_mcp_workbook_discovery_cells_and_cached_reuse(
    running, bridge, owner, seeded, readable_workbook, token_factory, tmp_path, monkeypatch
):
    from canvas_bridge.parser_sandbox import extract_document

    _, ref = readable_workbook
    bridge.settings.parser_root = tmp_path / "parser"
    monkeypatch.setattr("canvas_bridge.course_files.extract_document", extract_document)
    async with httpx.AsyncClient(
        headers={"Authorization": f"Bearer {token_factory(scope='academic:read')}"}, timeout=60
    ) as http:
        async with streamable_http_client(running + "/mcp", http_client=http) as streams:
            async with ClientSession(streams[0], streams[1]) as session:
                await session.initialize()
                tool = next(
                    t
                    for t in (await session.list_tools()).tools
                    if t.name == "read_course_workbook"
                )
                assert tool.annotations.readOnlyHint and not tool.annotations.destructiveHint
                response = await session.call_tool(
                    "read_course_workbook", {"connection_id": seeded, "file_id": ref}
                )
                assert response.structuredContent["status"] == "ok"
                assert [
                    s["name"] for s in response.structuredContent["data"]["workbook"]["worksheets"]
                ] == ["Analysis", "Reference", "Empty"]
                response = await session.call_tool(
                    "read_course_workbook",
                    {
                        "connection_id": seeded,
                        "file_id": ref,
                        "sheet_name": "Analysis",
                        "cell_range": "D2:D3",
                    },
                )
                result = response.structuredContent
                assert result["status"] == "ok"
                assert result["data"]["cells"][0]["cached_value"] == 5.3
                assert result["data"]["cells"][1]["cached_value_status"] == "unavailable"
                assert result["sources"][0]["locator"] == "'Analysis'!D2"
                assert result["sources"][0]["content_hash"] == result["data"]["content_hash"]
    assert sum("download" in r.url.path for r in bridge.synthetic.requests) == 1
    assert all(r.method == "GET" for r in bridge.synthetic.requests)
