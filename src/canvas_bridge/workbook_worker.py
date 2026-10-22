"""Read-only bounded OOXML workbook extraction inside the existing OS sandbox."""

import io
import json
import math
from datetime import date, datetime, time, timedelta

if not globals().get("__package__"):
    # openpyxl constructs a MIME table at import time. In the isolated worker,
    # use built-in mappings instead of probing host /etc/mime.types files.
    # This runs before any document is opened and adds no filesystem access.
    import mimetypes

    mimetypes.knownfiles = []
    mimetypes.init(files=[])

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

if globals().get("__package__"):
    from .office_package import OfficePackage
else:
    from office_package import OfficePackage

MAX_CELLS = 5000
MAX_ROWS = 5000
MAX_COLS = 100
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
WORKBOOK_ERRORS = {
    "workbook_cell_limit",
    "workbook_sheet_limit",
    "invalid_workbook_reference",
    "workbook_metadata_limit",
}


def scalar(value):
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return {"seconds": value.total_seconds()}
    if value is None or type(value) in {str, int, float, bool}:
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError("workbook_cell_limit")
        if isinstance(value, str) and len(value) > 8192:
            raise ValueError("workbook_cell_limit")
        return value
    return {"unsupported_type": type(value).__name__}


def extract_workbook(data):
    warnings = [
        "Formulas are source expressions, never evaluated. Cached results may be missing or stale; no recalculation was performed.",
        "Sheet visibility, number formats and merged regions are source metadata. Nearby labels are context, not verified scientific meanings.",
        "Charts, images, pivots, conditional formatting and external data refresh are not interpreted. External links are never followed.",
    ]
    with OfficePackage(data, "xlsx") as package:
        names, xml = package.names, package.xml
        # Validate XML before handing the package to a library; no DTD/entity fallback.
        for name in names:
            if name.endswith((".xml", ".rels")):
                xml(name)
        workbook_root = xml("xl/workbook.xml")
        sheets = list(workbook_root.find(NS + "sheets"))
        if not 1 <= len(sheets) <= 50:
            raise ValueError("workbook_sheet_limit")
        sheet_names = [s.attrib.get("name", "") for s in sheets]
        if any(not n or len(n) > 31 for n in sheet_names) or len(
            set(n.casefold() for n in sheet_names)
        ) != len(sheets):
            raise ValueError("invalid_workbook_reference")
        sheet_paths = []
        for sheet in sheets:
            relation = next((v for k, v in sheet.attrib.items() if k.endswith("}id")), None)
            sheet_paths.append(package.related_part("xl/workbook.xml", relation, "worksheet"))
        metadata = {
            "worksheet_count": len(sheets),
            "worksheets": [],
            "formula_evaluation": "not_performed",
            "external_features_present": any(
                "externalLinks/" in n or n == "xl/connections.xml" for n in names
            ),
        }
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        cached = None
        try:
            cached = load_workbook(
                io.BytesIO(data), read_only=True, data_only=True, keep_links=False
            )
            metadata["date_system"] = "1904" if wb.epoch.year == 1904 else "1900"
            metadata["named_ranges"] = [
                {
                    "name": n.attrib.get("name"),
                    "reference": n.text,
                    "local_sheet_id": n.attrib.get("localSheetId"),
                }
                for n in workbook_root.iter(NS + "definedName")
            ]
            pages, total, text_total, structured_total = [], 0, 0, 0
            partial = False
            for index, sheet in enumerate(sheets, 1):
                sheet_name = sheet.attrib["name"]
                path = sheet_paths[index - 1]
                root = xml(path)
                raw_entries = list(root.iter(NS + "c"))
                raw_cells = {e.attrib["r"]: e for e in raw_entries if "r" in e.attrib}
                if len(raw_cells) != len(raw_entries):
                    raise ValueError("invalid_workbook_reference")
                if len(raw_cells) > 100_000:
                    raise ValueError("workbook_cell_limit")
                dimensions = root.find(NS + "dimension")
                declared = dimensions.get("ref") if dimensions is not None else None
                merged = [e.attrib["ref"] for e in root.iter(NS + "mergeCell")]
                all_bounds = [range_boundaries(a + ":" + a) for a in raw_cells]
                if any(not (1 <= b[0] <= 16384 and 1 <= b[1] <= 1048576) for b in all_bounds):
                    raise ValueError("invalid_workbook_reference")
                max_col = max((b[2] for b in all_bounds), default=0)
                max_row = max((b[3] for b in all_bounds), default=0)
                used = (
                    f"{get_column_letter(min(b[0] for b in all_bounds))}{min(b[1] for b in all_bounds)}:{get_column_letter(max_col)}{max_row}"
                    if all_bounds
                    else None
                )
                tables = []
                for table_part in root.iter(NS + "tablePart"):
                    relation = next(
                        (v for k, v in table_part.attrib.items() if k.endswith("}id")), None
                    )
                    table = xml(package.related_part(path, relation, "table"))
                    tables.append(
                        {
                            "name": table.get("displayName"),
                            "range": table.get("ref"),
                            "headers": [c.get("name") for c in table.iter(NS + "tableColumn")],
                        }
                    )
                info = {
                    "name": sheet_name,
                    "index": index,
                    "visibility": sheet.get("state", "visible"),
                    "declared_dimensions": declared,
                    "used_range": used,
                    "merged_regions": merged,
                    "tables": tables,
                    "stored_cell_count": len(raw_cells),
                    "nonempty_cell_count": sum(
                        any(e.find(NS + k) is not None for k in ("v", "is", "f"))
                        for e in raw_cells.values()
                    ),
                    "formula_cell_count": sum(
                        e.find(NS + "f") is not None for e in raw_cells.values()
                    ),
                    "coverage": "complete",
                    "extracted_cell_count": 0,
                }
                metadata["worksheets"].append(info)
                if len(json.dumps(metadata, ensure_ascii=True)) > 30000:
                    raise ValueError("workbook_metadata_limit")
                ws, cs = wb[sheet_name], cached[sheet_name]
                # Ignore potentially false declared dimensions, but never walk unbounded sparse rows.
                ws.reset_dimensions()
                cs.reset_dimensions()
                records = []
                limited = False
                if max_row and max_col:
                    bounds = {"max_row": min(max_row, MAX_ROWS), "max_col": min(max_col, MAX_COLS)}
                    for row, crow in zip(
                        ws.iter_rows(**bounds), cs.iter_rows(**bounds), strict=True
                    ):
                        for cell, cache in zip(row, crow, strict=True):
                            if cell.value is None:
                                continue
                            if total >= MAX_CELLS:
                                break
                            address = cell.coordinate
                            raw = raw_cells.get(address)
                            rv = raw.find(NS + "v") if raw is not None else None
                            f = raw.find(NS + "f") if raw is not None else None
                            if f is not None:
                                scalar(f.text)
                                for value in f.attrib.values():
                                    scalar(value)
                            formula = scalar(cell.value) if cell.data_type == "f" else None
                            if f is not None and not isinstance(formula, str):
                                formula = {
                                    "text": f.text,
                                    "attributes": f.attrib,
                                    "status": "unsupported_formula_type",
                                }
                            record = {
                                "sheet": sheet_name,
                                "address": address,
                                "row": cell.row,
                                "column": cell.column,
                                "value": None if formula is not None else scalar(cell.value),
                                "raw_value": rv.text
                                if rv is not None
                                else scalar(cell.value)
                                if formula is None
                                else None,
                                "data_type": cell.data_type,
                                "formula": formula,
                                "cached_value": scalar(cache.value)
                                if formula is not None
                                else None,
                                "cached_value_status": (
                                    "available_unverified_freshness"
                                    if rv is not None and rv.text is not None
                                    else "unavailable"
                                )
                                if formula is not None
                                else "not_applicable",
                                "number_format": scalar(cell.number_format),
                                "date_time_interpretation": scalar(cell.value)
                                if cell.is_date
                                else None,
                                "error": cell.value
                                if cell.data_type == "e"
                                else cache.value
                                if cache.data_type == "e"
                                else None,
                                "cross_sheet_reference": bool(formula and "!" in str(formula)),
                                "possible_external_reference": bool(
                                    formula and "[" in str(formula) and "!" in str(formula)
                                ),
                                "reference_detection": "lexical_unverified",
                            }
                            line = f"'{sheet_name.replace(chr(39), chr(39) * 2)}'!{address}: {json.dumps({'value': record['value'], 'formula': formula, 'cached_value': record['cached_value'], 'cached_value_status': record['cached_value_status']}, ensure_ascii=True)}"
                            # Count separators as well as text. Keep the cached structured
                            # view within the worker's 2 MB JSON result budget too.
                            record_size = len(json.dumps(record, ensure_ascii=True))
                            if (
                                text_total + len(line) + 2 > 240_000
                                or structured_total + record_size > 900_000
                            ):
                                limited = True
                                break
                            text_total += len(line) + 2
                            structured_total += record_size
                            record["index_text"] = line
                            records.append(record)
                            total += 1
                        if total >= MAX_CELLS or limited:
                            break
                info["extracted_cell_count"] = len(records)
                if (
                    max_row > MAX_ROWS
                    or max_col > MAX_COLS
                    or len(records) < info["nonempty_cell_count"]
                ):
                    info["coverage"] = "partial"
                    partial = True
                for at in range(0, max(1, len(records)), 100):
                    cells = records[at : at + 100]
                    text = "\n".join(c.pop("index_text") for c in cells)
                    if not text:
                        text = f"Worksheet {sheet_name!r}: {info['coverage']} cell coverage; no extracted nonempty cells."
                    pages.append(
                        {
                            "page": len(pages) + 1,
                            "text": text,
                            "state": "extracted" if text else "no_text",
                            "locator_kind": "worksheet_range",
                            "sheet": sheet_name,
                            "cells": cells,
                            "range": f"{get_column_letter(min(c['column'] for c in cells))}{min(c['row'] for c in cells)}:{get_column_letter(max(c['column'] for c in cells))}{max(c['row'] for c in cells)}"
                            if cells
                            else None,
                            "extraction_method": "xlsx_stored_cells",
                        }
                    )
            if partial:
                warnings.append(
                    "Workbook extraction is partial: at most 5000 nonempty cells, rows 1–5000, columns A–CV, 240000 cell-index characters and 900000 structured-cell JSON characters. Unreturned cells are not evidence of emptiness."
                )
            return metadata, pages, partial, warnings
        finally:
            wb.close()
            if cached:
                cached.close()
