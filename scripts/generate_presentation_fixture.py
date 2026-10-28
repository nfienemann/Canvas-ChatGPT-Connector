"""Generate a fictional image-heavy rendering fixture using python-pptx.

Development only; no parser runtime dependency. Run with the bundled artifact
Python. This creates only synthetic content in tests/fixtures.
"""

import io
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def main():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(5.625)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title.text = "SYNTHETIC ONLY: fluorescence controls"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(10, 30, 60)
    figure = Image.new("RGB", (900, 380), "#eeeeee")
    draw = ImageDraw.Draw(figure)
    draw.rectangle((50, 50, 380, 330), fill="#11334b")
    draw.rectangle((500, 50, 830, 330), fill="#11334b")
    for x, y in [(100, 100), (180, 240), (310, 140)]:
        draw.ellipse((x, y, x + 25, y + 25), fill="#47da97")
    for x, y in [(550, 100), (630, 240), (760, 140), (550, 240), (630, 100), (760, 240)]:
        draw.ellipse((x, y, x + 25, y + 25), fill="#47da97")
    draw.text((140, 350), "CONTROL", fill="black", font_size=22)
    draw.text((600, 350), "TREATMENT", fill="black", font_size=22)
    image = io.BytesIO()
    figure.save(image, format="PNG")
    slide.shapes.add_picture(image, Inches(0.5), Inches(1.1), width=Inches(9))
    slide.notes_slide.notes_text_frame.text = (
        "Fictional diagram; spot counts alone cannot establish efficacy."
    )
    image_only = deck.slides.add_slide(deck.slide_layouts[6])
    image_only.shapes.add_picture(
        io.BytesIO(image.getvalue()), 0, 0, width=deck.slide_width, height=deck.slide_height
    )
    blank = deck.slides.add_slide(deck.slide_layouts[6])
    blank._element.set("show", "0")
    target = Path(__file__).resolve().parents[1] / "tests/fixtures/synthetic_visual_slides.pptx"
    target.parent.mkdir(exist_ok=True)
    deck.save(target)
    print("Generated synthetic_visual_slides.pptx (three fictional slides)")


if __name__ == "__main__":
    main()
